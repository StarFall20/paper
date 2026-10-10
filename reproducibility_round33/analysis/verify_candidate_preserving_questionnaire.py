"""Audit the respondent-facing candidate-preserving questionnaire files.

The earlier audit reconstructed the cards from constants in this module. That
could prove the constants internally consistent while missing an editing error
in the actual DOCX files. This version reads every table from each DOCX, maps
the displayed levels back to preregistered codes, and checks the candidate
vector, shared attribute, and price card.
"""
from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path

from docx import Document


TASKS = [
    ("G0", "G1", 79, 99),
    ("G2", "G1", 59, 89),
    ("G1", "G0", 69, 109),
    ("G0", "G2", 79, 119),
    ("G2", "G1", 89, 69),
    ("G1", "G0", 99, 79),
    ("G0", "G2", 109, 59),
    ("G2", "G1", 119, 69),
]
V1 = [("D0", "S2", "I0", "E2"), ("D1", "S2", "I0", "E1")]
V2 = [("D1", "S1", "I1", "E1"), ("D2", "S1", "I1", "E0")]
LEVEL = {
    "D0": 0, "D1": 1, "D2": 2,
    "S0": 0, "S1": 1, "S2": 2,
    "I0": 0, "I1": 1, "I2": 2,
    "E0": 0, "E1": 1, "E2": 2,
}

LABELS = {
    "cn": {
        "数据管理": {"仅在设备上保存": "D0", "本地保存并加密备份": "D1", "在安全云端处理": "D2"},
        "专业支持": {"无专业支持": "S0", "护士在线咨询": "S1", "专业团队支持": "S2"},
        "智能功能": {"基础提醒": "I0", "自适应提醒": "I1", "个性化建议": "I2"},
        "临床证据": {"无临床证据": "E0", "初步临床证据": "E1", "充分临床证据": "E2"},
        "数据共享": {"不共享数据": "G0", "匿名科研共享": "G1", "经同意后可控共享": "G2"},
        "月费": None,
    },
    "en": {
        "Data management": {"On-device only": "D0", "Local storage with encrypted backup": "D1", "Processed in a secure cloud": "D2", "Secure cloud processing": "D2"},
        "Professional support": {"No professional support": "S0", "Nurse consultation": "S1", "Nurse chat": "S1", "Specialist team support": "S2"},
        "Smart functions": {"Basic reminders": "I0", "Adaptive reminders": "I1", "Adaptive alerts": "I1", "Personalized recommendations": "I2"},
        "Clinical evidence": {"No clinical evidence": "E0", "Preliminary clinical evidence": "E1", "Strong clinical evidence": "E2"},
        "Data sharing": {"No data sharing": "G0", "Anonymous research sharing": "G1", "Controlled sharing with your approval": "G2"},
        "Monthly fee": None,
    },
}

ATTRIBUTE_KEY = {
    "数据管理": "data", "Data management": "data",
    "专业支持": "support", "Professional support": "support",
    "智能功能": "intelligence", "Smart functions": "intelligence",
    "临床证据": "evidence", "Clinical evidence": "evidence",
    "数据共享": "sharing", "Data sharing": "sharing",
    "月费": "price", "Monthly fee": "price",
}


def vector(profile):
    data, support, intelligence, evidence = profile
    return (LEVEL[data] + LEVEL[support], LEVEL[intelligence] + LEVEL[evidence])


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def _parse_price(text: str) -> int:
    match = re.search(r"\d+", text.replace(",", ""))
    if not match:
        raise ValueError(f"price cell has no integer: {text!r}")
    return int(match.group(0))


def read_docx(path: str | Path):
    """Return eight parsed task cards and the detected language."""
    path = Path(path)
    doc = Document(path)
    if len(doc.tables) != len(TASKS):
        raise ValueError(f"{path}: expected {len(TASKS)} task tables, found {len(doc.tables)}")
    first_header = _clean(doc.tables[0].rows[0].cells[0].text)
    language = "cn" if first_header == "属性" else "en" if first_header == "Attribute" else None
    if language is None:
        raise ValueError(f"{path}: unrecognised table header {first_header!r}")
    label_map = LABELS[language]
    cards = []
    for task_no, table in enumerate(doc.tables, 1):
        if len(table.rows) != 7 or len(table.columns) != 3:
            raise ValueError(f"{path}: task {task_no} expected a 7-by-3 table")
        header = [_clean(cell.text) for cell in table.rows[0].cells]
        expected_header = ["属性", "服务 A", "服务 B"] if language == "cn" else ["Attribute", "Service A", "Service B"]
        if header != expected_header:
            raise ValueError(f"{path}: task {task_no} header {header!r} != {expected_header!r}")
        parsed = {"data": [], "support": [], "intelligence": [], "evidence": [], "sharing": [], "price": []}
        for row in table.rows[1:]:
            cells = [_clean(cell.text) for cell in row.cells]
            attr = cells[0]
            if attr not in label_map:
                raise ValueError(f"{path}: task {task_no} unknown attribute {attr!r}")
            key = ATTRIBUTE_KEY[attr]
            if key == "price":
                parsed[key] = [_parse_price(cells[1]), _parse_price(cells[2])]
            else:
                try:
                    parsed[key] = [label_map[attr][cells[1]], label_map[attr][cells[2]]]
                except KeyError as exc:
                    raise ValueError(f"{path}: task {task_no} unknown {attr} level {exc.args[0]!r}") from exc
        if any(len(parsed[key]) != 2 for key in parsed):
            raise ValueError(f"{path}: task {task_no} is missing an attribute row")
        cards.append(parsed)
    return language, cards


def audit_pair(v1_path: str | Path, v2_path: str | Path):
    lang1, cards1 = read_docx(v1_path)
    lang2, cards2 = read_docx(v2_path)
    if lang1 != lang2:
        raise ValueError(f"language mismatch: {v1_path}={lang1}, {v2_path}={lang2}")
    rows = []
    for task, (card1, card2, (ga, gb, pa, pb)) in enumerate(zip(cards1, cards2, TASKS), 1):
        expected_sharing = [ga, gb]
        expected_prices = [pa, pb]
        for alt in range(2):
            p1 = (card1["data"][alt], card1["support"][alt], card1["intelligence"][alt], card1["evidence"][alt])
            p2 = (card2["data"][alt], card2["support"][alt], card2["intelligence"][alt], card2["evidence"][alt])
            expected_p1 = V1[alt]
            expected_p2 = V2[alt]
            sharing_fixed = card1["sharing"][alt] == card2["sharing"][alt] == expected_sharing[alt]
            price_fixed = card1["price"][alt] == card2["price"][alt] == expected_prices[alt]
            rows.append({
                "language": lang1,
                "task": task,
                "alternative": alt + 1,
                "v1_profile": ":".join(p1),
                "v2_profile": ":".join(p2),
                "profile_layout_matches_spec": int(p1 == expected_p1 and p2 == expected_p2),
                "v1_candidate_1": vector(p1)[0],
                "v2_candidate_1": vector(p2)[0],
                "v1_candidate_2": vector(p1)[1],
                "v2_candidate_2": vector(p2)[1],
                "v1_sharing": card1["sharing"][alt],
                "v2_sharing": card2["sharing"][alt],
                "v1_price": card1["price"][alt],
                "v2_price": card2["price"][alt],
                "sharing_and_price_fixed": int(sharing_fixed and price_fixed),
                "candidate_vector_preserved": int(vector(p1) == vector(p2)),
            })
    if not all(row["sharing_and_price_fixed"] and row["candidate_vector_preserved"] and row["profile_layout_matches_spec"] for row in rows):
        raise AssertionError("DOCX pair failed profile, candidate, sharing, or price preservation")
    return rows


def run(v1="survey/candidate_preserving_dce_questionnaire_v1.docx",
        v2="survey/candidate_preserving_dce_questionnaire_v2.docx",
        out="results/questionnaire_candidate_preservation_audit.csv",
        english_v1=None, english_v2=None):
    pairs = [(v1, v2)]
    if (english_v1 is None) != (english_v2 is None):
        raise ValueError("provide both --english-v1 and --english-v2, or neither")
    if english_v1:
        pairs.append((english_v1, english_v2))
    rows = []
    for first, second in pairs:
        rows.extend(audit_pair(first, second))
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    with open(out, "w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    n_languages = len(set(row["language"] for row in rows))
    print(f"verified {len(rows)} rendered DOCX cards across 8 tasks and {n_languages} language pair(s); all checks passed")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--v1", default="survey/candidate_preserving_dce_questionnaire_v1.docx")
    parser.add_argument("--v2", default="survey/candidate_preserving_dce_questionnaire_v2.docx")
    parser.add_argument("--english-v1")
    parser.add_argument("--english-v2")
    parser.add_argument("--out", default="results/questionnaire_candidate_preservation_audit.csv")
    args = parser.parse_args()
    run(**vars(args))
